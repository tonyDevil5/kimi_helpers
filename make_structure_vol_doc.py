"""Word summary of daily realized vol for CC basis, coupon swaps and coupon flys."""
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt

from mbs_rv_report import (
    COUPON_FLY_CENTERS,
    COUPON_SWAP_PAIRS,
    STRUCTURE_AGENCIES,
    build_coupon_fly_price_timeseries,
    build_coupon_swap_jpm_perf_timeseries,
    load_mbs_rv,
)

W = [20, 40, 60]
summary = pd.read_csv("mbs_rv_output/structure_daily_vol_summary.csv")
as_of = "2026-09-21"

# Composite (equal-weighted average daily change) vols
dat = load_mbs_rv()
swap = build_coupon_swap_jpm_perf_timeseries(dat)
fly = build_coupon_fly_price_timeseries(dat)
sk = lambda ag, p: f"{ag.lower()}_{p[0]:g}_{p[1]:g}_swap_perf_32nds"
fk = lambda ag, c: f"{ag.lower()}_{c:g}_fly_change_32nds"

comp_rows = []
for group, df, key_fn, items, label in [
    ("cpn_swap", swap, sk, COUPON_SWAP_PAIRS, "Coupon swaps"),
    ("fly", fly, fk, COUPON_FLY_CENTERS, "Coupon flys"),
]:
    per_ag = {ag: [df[key_fn(ag, x)] for x in items] for ag in STRUCTURE_AGENCIES}
    comp = {ag: pd.concat(c, axis=1).mean(axis=1) for ag, c in per_ag.items()}
    comp["ALL"] = pd.concat([s for c in per_ag.values() for s in c], axis=1).mean(axis=1)
    indiv = dict(per_ag)
    indiv["ALL"] = [s for c in per_ag.values() for s in c]
    for name, s in comp.items():
        s = s.dropna()
        vols, means = [], []
        vols.append(s.std())
        means.append(pd.Series([x.std() for x in indiv[name]]).mean())
        for w in W:
            vols.append(s.rolling(w, min_periods=w).std().iloc[-1])
            means.append(pd.Series(
                [x.rolling(w, min_periods=w).std().iloc[-1] for x in indiv[name]]).mean())
        comp_rows.append((label, name, vols, means))

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10)

doc.add_heading("MBS Structure Daily Realized Vol", level=0)
doc.add_paragraph(
    f"Daily realized vol (standard deviation of daily changes) over the full sample "
    f"(2023-09-19 onward) and trailing 20/40/60-day windows, as of {as_of}. CC basis is in bps (std of daily OAS changes, Bloomberg "
    ".FNCC105 / .G2CC105); coupon swaps and coupon flys are in ticks (32nds) - the JPM "
    "perf-swap and price-fly series are already daily changes, so vol is the std of the "
    "series itself. Source: mbs_rv.xlsx / Bloomberg, as used in mbs_rv_report.py."
)


def add_table(headers, rows, fmt="{:.2f}"):
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            t.rows[i].cells[j].text = v if isinstance(v, str) else fmt.format(v)
    return t


def vol_rows(group):
    out = []
    for _, r in summary[summary.group == group].iterrows():
        out.append([r["structure"], r["vol_all"], r["vol_20d"], r["vol_40d"], r["vol_60d"]])
    return out


VOL_HEADERS = ["vol full", "vol 20d", "vol 40d", "vol 60d"]

doc.add_heading("1. Individual Structures", level=1)

doc.add_heading("CC basis (bps)", level=2)
add_table(["Series", *VOL_HEADERS], vol_rows("cc_basis"))

doc.add_heading("Coupon swaps - JPM adjacent-coupon perf (ticks)", level=2)
add_table(["Swap", *VOL_HEADERS], vol_rows("cpn_swap"))

doc.add_heading("Coupon flys (ticks)", level=2)
add_table(["Fly", *VOL_HEADERS], vol_rows("fly"))

doc.add_heading("2. Averaged Across Structures", level=1)
doc.add_paragraph(
    '"Composite vol" is the vol of the equal-weighted average daily change across the '
    "structures in the group (diversification across structures lowers it); "
    '"mean indiv vol" is the simple average of the individual vols.'
)
add_table(
    ["Group", "Avg over", "comp full", "comp 20d", "comp 40d", "comp 60d",
     "mean indiv full", "mean indiv 20d", "mean indiv 40d", "mean indiv 60d"],
    [[label, name, *vols, *means] for label, name, vols, means in comp_rows],
)
doc.add_paragraph()
doc.add_paragraph(
    "G2 structures run roughly 1.5-2x the FN vol in both swaps and flys. The composite "
    "vol is only ~0.4-0.5 ticks because daily moves across coupon pairs/flys partially "
    "offset each other."
)

import datetime

out_path = r"D:\python\kimi_helpers\MBS_Structure_Daily_Vol_Summary.docx"
try:
    doc.save(out_path)
except PermissionError:
    stamp = datetime.date.today().strftime("%Y%m%d")
    out_path = rf"D:\python\kimi_helpers\MBS_Structure_Daily_Vol_Summary_{stamp}.docx"
    doc.save(out_path)
    print("Primary file locked (open in Word?); saved dated copy instead.")
print(f"Saved {out_path}")
