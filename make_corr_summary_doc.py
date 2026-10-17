"""Summarize daily-change correlations to CT10 (full sample + past 120d) in a Word doc."""
from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10)

doc.add_heading("MBS Daily-Change Correlations to CT10", level=0)

doc.add_paragraph(
    "Correlations of daily changes versus daily changes in the 10-year Treasury yield "
    "(CT10). Full sample: 2023-09-19 to 2026-09-21 (748 daily observations); the 120d "
    "column is the rolling 120-day correlation as of 2026-09-21. Source: mbs_rv.xlsx / "
    "Bloomberg current-coupon OAS, as used in mbs_rv_report.py."
)


def add_table(headers, rows):
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            t.rows[i].cells[j].text = v if isinstance(v, str) else f"{v:+.2f}"
    return t


doc.add_heading("1. FN / G2 Current-Coupon Basis (OAS)", level=1)
doc.add_paragraph("Daily change in current-coupon OAS vs daily change in CT10.")
add_table(
    ["Series", "Full-sample corr", "120d corr"],
    [
        ["FN CC basis (FNCC)", 0.39, 0.48],
        ["G2 CC basis (G2CC)", 0.37, 0.35],
    ],
)
doc.add_paragraph()
doc.add_paragraph(
    "Both bases widen when yields rise. The level correlation over the full sample is "
    "much lower (FN +0.16, G2 +0.11): yield levels do not pin down the basis level, but "
    "day-to-day yield moves do drive day-to-day basis moves."
)

doc.add_heading("2. Perf-Based Adjacent Coupon Swaps", level=1)
doc.add_paragraph(
    "Daily performance differential (upper coupon minus lower coupon, JPM Perf in 32nds) "
    "vs daily CT10 change."
)
add_table(
    ["Swap", "FN full sample", "FN 120d", "G2 full sample", "G2 120d"],
    [
        ["3.5/3", -0.20, -0.23, -0.07, +0.02],
        ["4/3.5", -0.11, -0.11, -0.12, -0.02],
        ["4.5/4", +0.06, +0.02, +0.05, -0.02],
        ["5/4.5", +0.28, -0.05, +0.29, -0.28],
        ["5.5/5", +0.41, +0.54, +0.47, +0.68],
        ["6/5.5", +0.49, +0.48, +0.28, +0.16],
        ["6.5/6", +0.41, -0.09, +0.27, +0.09],
    ],
)
doc.add_paragraph()
doc.add_paragraph(
    "Clear coupon-stack pattern: below the 5% coupon the swap is negatively correlated "
    "with yield changes (lower coupons outperform on selloffs); above 5% the correlation "
    "turns strongly positive (up-in-coupon outperforms when yields back up), consistent "
    "with convexity / prepayment dynamics. Over the past 120 days the pattern is intact "
    "at 5.5/5 and 6/5.5 (G2 5.5/5 up to +0.68), while FN 6.5/6 and G2 5/4.5 have "
    "recently decoupled from their full-sample relationships."
)

doc.add_heading("3. Coupon Flys", level=1)
doc.add_paragraph(
    "Daily change of the price fly level (2 x center - wings, in 32nds) vs daily CT10 "
    "change. Note: the report's fly series is price-based, not perf-based."
)
add_table(
    ["Fly", "FN full sample", "FN 120d", "G2 full sample", "G2 120d"],
    [
        ["3.5 fly", -0.09, +0.01, -0.07, -0.08],
        ["4 fly", -0.34, -0.43, -0.19, -0.35],
        ["4.5 fly", -0.44, -0.59, -0.49, -0.60],
        ["5 fly", -0.57, -0.78, -0.32, -0.52],
        ["5.5 fly", -0.18, -0.36, +0.06, +0.17],
        ["6 fly", +0.35, +0.39, +0.21, +0.27],
    ],
)
doc.add_paragraph()
doc.add_paragraph(
    "Mid-coupon flys (4-5) are strongly negatively correlated: the belly underperforms "
    "the wings on selloffs; the 6 fly is positive. The past 120 days show the same "
    "structure with even stronger magnitudes (FN 5 fly -0.78, G2 4.5 fly -0.60)."
)

doc.add_heading("4. G2/FN Perf Swaps", level=1)
doc.add_paragraph(
    "Daily G2-minus-FN performance (JPM Perf, 32nds) vs daily CT10 change."
)
add_table(
    ["Coupon", "Full-sample corr", "120d corr"],
    [
        ["G2/FN 3", -0.04, -0.07],
        ["G2/FN 3.5", +0.02, +0.13],
        ["G2/FN 4", -0.06, +0.26],
        ["G2/FN 4.5", -0.07, +0.18],
        ["G2/FN 5", +0.08, -0.15],
        ["G2/FN 5.5", +0.31, +0.47],
        ["G2/FN 6", +0.19, +0.16],
        ["G2/FN 6.5", +0.22, +0.29],
    ],
)
doc.add_paragraph()
doc.add_paragraph(
    "Largely rate-neutral except at high coupons (5.5-6.5), where FN's worse convexity "
    "makes G2 outperform on selloffs. The 120d window shows the same high-coupon tilt, "
    "with G2/FN 5.5 strengthening to +0.47."
)

doc.add_heading("5. Level vs Daily-Change Correlation", level=1)
doc.add_paragraph(
    "Level correlation uses the raw series and mixes long-run regime effects with "
    "short-run linkage; for yields and OAS it is typically weak. Daily-change "
    "correlation differences out the trend and measures how the series react to the same "
    "shock today. Here the basis/swap/fly levels have their own slow-moving drivers, "
    "while daily moves are strongly yield-linked - so daily-change correlations are the "
    "actionable numbers for hedging and risk management."
)

doc.save(r"D:\python\kimi_helpers\MBS_CT10_Correlation_Summary.docx")
print("Saved MBS_CT10_Correlation_Summary.docx")
